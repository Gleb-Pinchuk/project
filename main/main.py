import json
import csv
from openpyxl import load_workbook
from datetime import datetime
from typing import List, Dict
from src.process_bank import process_bank_search


AVAILABLE_STATUSES = ["EXECUTED", "CANCELED", "PENDING"] 


def load_from_json(filepath: str) -> List[Dict]:
    with open(filepath, "r", encoding="utf-8") as f:
        return json.load(f)


def load_from_csv(filepath: str) -> List[Dict]:
    with open(filepath, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader)


def load_from_xlsx(filepath: str) -> List[Dict]:
    wb = load_workbook(filepath)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        data.append(dict(zip(headers, row)))
    return data


def print_operations(operations: List[Dict]) -> None:
    if not operations:
        print("\nПрограмма: Не найдено ни одной транзакции, подходящей под ваши условия фильтрации")
        return

    print("\nПрограмма: Распечатываю итоговый список транзакций...")
    print(f"\nВсего банковских операций в выборке: {len(operations)}\n")

    for op in operations:
        date = datetime.strptime(op["date"], "%Y-%m-%d").strftime("%d.%m.%Y")
        print(f"{date} {op['description']}")
        print(op.get("from", ""), "->", op.get("to", "")) if op.get("from") else print(op.get("to", ""))
        print(f"Сумма: {op['amount']} {op['currency']}\n")


def main():
    print("Программа: Привет! Добро пожаловать в программу работы с банковскими транзакциями.")
    print("Выберите необходимый пункт меню:")
    print("1. Получить информацию о транзакциях из JSON-файла")
    print("2. Получить информацию о транзакциях из CSV-файла")
    print("3. Получить информацию о транзакциях из XLSX-файла")

    choice = input("\nПользователь: ").strip()

    if choice == "1":
        print("\nПрограмма: Для обработки выбран JSON-файл.")
        data = load_from_json("data/operations.json")
    elif choice == "2":
        print("\nПрограмма: Для обработки выбран CSV-файл.")
        data = load_from_csv("data/operations.csv")
    elif choice == "3":
        print("\nПрограмма: Для обработки выбран XLSX-файл.")
        data = load_from_xlsx("data/operations.xlsx")
    else:
        print("\nПрограмма: Некорректный выбор. Завершение работы.")
        return

    # фильтрация по статусу
    while True:
        status = input(
            '\nПрограмма: Введите статус, по которому необходимо выполнить фильтрацию. '
            'Доступные для фильтровки статусы: EXECUTED, CANCELED, PENDING\n\nПользователь: '
        ).strip().upper()

        if status in AVAILABLE_STATUSES:
            print(f'\nПрограмма: Операции отфильтрованы по статусу "{status}"')
            data = [op for op in data if op.get("status", "").upper() == status]
            break
        else:
            print(f'\nПрограмма: Статус операции "{status}" недоступен.')

    # сортировка по дате
    sort_choice = input("\nПрограмма: Отсортировать операции по дате? Да/Нет\n\nПользователь: ").strip().lower()
    if sort_choice == "да":
        order = input("\nПрограмма: Отсортировать по возрастанию или по убыванию?\n\nПользователь: ").strip().lower()
        reverse = True if order == "по убыванию" else False
        data.sort(key=lambda x: datetime.strptime(x["date"], "%Y-%m-%d"), reverse=reverse)

    # фильтрация по рублевым транзакциям
    rub_choice = input("\nПрограмма: Выводить только рублевые транзакции? Да/Нет\n\nПользователь: ").strip().lower()
    if rub_choice == "да":
        data = [op for op in data if op.get("currency", "").upper() in ("RUB", "РУБ", "РУБЛЬ")]

    # фильтрация по слову в описании
    desc_choice = input("\nПрограмма: Отфильтровать список транзакций по определенному слову в описании? Да/Нет\n\nПользователь: ").strip().lower()
    if desc_choice == "да":
        word = input("\nПрограмма: Введите слово для поиска в описании\n\nПользователь: ").strip()
        data = process_bank_search(data, word)
